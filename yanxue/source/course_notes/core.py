from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import subprocess
import time
import tomllib
import urllib.error
import urllib.request
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree


MEDIA_EXTENSIONS = {".mp4", ".m4a", ".mp3", ".wav", ".webm", ".mpeg", ".mpga"}
DOCUMENT_EXTENSIONS = {".pdf", ".pptx", ".docx", ".txt", ".md"}
IGNORED_NAMES = {"draft.md", "notes.md", "qa_report.md", "evidence.md"}


class PipelineError(RuntimeError):
    pass


def read_toml(path: Path) -> dict[str, Any]:
    with path.open("rb") as handle:
        return tomllib.load(handle)


def atomic_write(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(content, encoding="utf-8")
    temporary.replace(path)


def write_json(path: Path, data: Any) -> None:
    atomic_write(path, json.dumps(data, ensure_ascii=False, indent=2))


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def slugify(value: str) -> str:
    value = re.sub(r'[\\/:*?"<>|]', "_", value).strip().strip(".")
    value = re.sub(r"\s+", "_", value)
    return value[:100] or "未命名课程"


def seconds_to_clock(value: float) -> str:
    seconds = max(0, int(value))
    return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


def run_checked(command: list[str], cwd: Path | None = None) -> subprocess.CompletedProcess[str]:
    try:
        result = subprocess.run(
            command,
            cwd=cwd,
            text=True,
            encoding="utf-8",
            errors="replace",
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=False,
        )
    except FileNotFoundError as exc:
        raise PipelineError(f"未找到外部工具：{command[0]}") from exc
    if result.returncode:
        detail = (result.stderr or result.stdout)[-2000:]
        raise PipelineError(f"命令执行失败（{result.returncode}）：{' '.join(command)}\n{detail}")
    return result


@dataclass
class Evidence:
    source: str
    locator: str
    text: str
    kind: str

    def as_dict(self) -> dict[str, str]:
        return {"source": self.source, "locator": self.locator, "text": self.text, "kind": self.kind}


def extract_text_file(path: Path) -> list[Evidence]:
    text = path.read_text(encoding="utf-8", errors="replace")
    marker = re.compile(r"^(?:=====\s*)?(?:PDF )?PAGE\s+(\d+)(?:\s*=====)?$", re.I | re.M)
    matches = list(marker.finditer(text))
    if not matches:
        return [Evidence(path.name, "全文", text.strip(), "text")]
    output: list[Evidence] = []
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        body = text[start:end].strip()
        if body:
            output.append(Evidence(path.name, f"第 {match.group(1)} 页", body, "slide"))
    return output


def extract_pdf(path: Path, image_dir: Path) -> list[Evidence]:
    try:
        import pymupdf as fitz
    except ImportError as exc:
        try:
            import fitz  # type: ignore[no-redef]
        except ImportError:
            raise PipelineError("处理 PDF 需要 PyMuPDF：pip install PyMuPDF") from exc
    image_dir.mkdir(parents=True, exist_ok=True)
    document = fitz.open(path)
    output: list[Evidence] = []
    for number, page in enumerate(document, 1):
        text = page.get_text("text").strip()
        if text:
            output.append(Evidence(path.name, f"第 {number} 页", text, "slide"))
        preview = image_dir / f"{path.stem}_{number:04d}.png"
        if not preview.exists():
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False).save(preview)
    document.close()
    return output


def extract_pptx(path: Path) -> list[Evidence]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise PipelineError("处理 PPTX 需要 python-pptx：pip install python-pptx") from exc
    presentation = Presentation(path)
    output: list[Evidence] = []
    for number, slide in enumerate(presentation.slides, 1):
        blocks: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    blocks.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
        if blocks:
            output.append(Evidence(path.name, f"第 {number} 页", "\n".join(blocks), "slide"))
    return output


def extract_docx(path: Path) -> list[Evidence]:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    root = ElementTree.fromstring(xml)
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragraphs: list[str] = []
    for paragraph in root.iter(namespace + "p"):
        text = "".join(node.text or "" for node in paragraph.iter(namespace + "t")).strip()
        if text:
            paragraphs.append(text)
    return [Evidence(path.name, "全文", "\n\n".join(paragraphs), "transcript")]


def extract_document(path: Path, image_dir: Path) -> list[Evidence]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return extract_text_file(path)
    if suffix == ".pdf":
        return extract_pdf(path, image_dir)
    if suffix == ".pptx":
        return extract_pptx(path)
    if suffix == ".docx":
        return extract_docx(path)
    return []


def media_duration(path: Path) -> float:
    result = run_checked([
        "ffprobe", "-v", "error", "-show_entries", "format=duration",
        "-of", "default=noprint_wrappers=1:nokey=1", str(path),
    ])
    return float(result.stdout.strip())


def split_audio(path: Path, output_dir: Path, minutes: int = 12) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    pattern = output_dir / "chunk_%03d.mp3"
    expected = sorted(output_dir.glob("chunk_*.mp3"))
    if expected:
        return expected
    run_checked([
        "ffmpeg", "-hide_banner", "-loglevel", "error", "-i", str(path),
        "-vn", "-ac", "1", "-ar", "16000", "-b:a", "48k",
        "-f", "segment", "-segment_time", str(minutes * 60), "-reset_timestamps", "1",
        str(pattern),
    ])
    return sorted(output_dir.glob("chunk_*.mp3"))


class OpenAIProvider:
    def __init__(self, api_key: str | None = None, base_url: str | None = None):
        self.api_key = api_key or os.environ.get("OPENAI_API_KEY", "")
        self.base_url = (base_url or os.environ.get("OPENAI_BASE_URL") or "https://api.openai.com/v1").rstrip("/")
        if not self.api_key:
            raise PipelineError("没有检测到 OPENAI_API_KEY；可使用 --offline 只生成证据包。")

    def _request_json(self, endpoint: str, payload: dict[str, Any], timeout: int = 300) -> dict[str, Any]:
        request = urllib.request.Request(
            self.base_url + endpoint,
            data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
            headers={"Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json"},
            method="POST",
        )
        try:
            with urllib.request.urlopen(request, timeout=timeout) as response:
                return json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")
            raise PipelineError(f"OpenAI API 返回 {exc.code}：{detail[:1500]}") from exc

    def validate_key(self) -> None:
        request = urllib.request.Request(
            self.base_url + "/models",
            headers={"Authorization": f"Bearer {self.api_key}"},
            method="GET",
        )
        try:
            with urllib.request.urlopen(request, timeout=30) as response:
                if response.status != 200:
                    raise PipelineError(f"API 连通性检查失败：HTTP {response.status}")
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")
            raise PipelineError(f"API Key 无效或不可用（HTTP {exc.code}）：{detail[:500]}") from exc
        except urllib.error.URLError as exc:
            raise PipelineError(f"无法连接 API：{exc.reason}") from exc

    def structured_response(
        self,
        *,
        model: str,
        instructions: str,
        user_input: str,
        schema_name: str,
        schema: dict[str, Any],
    ) -> dict[str, Any]:
        payload = {
            "model": model,
            "instructions": instructions,
            "input": user_input,
            "text": {
                "format": {
                    "type": "json_schema",
                    "name": schema_name,
                    "strict": True,
                    "schema": schema,
                }
            },
        }
        response = self._request_json("/responses", payload)
        if response.get("status") not in {None, "completed"}:
            raise PipelineError(f"模型响应未完成：{response.get('status')} / {response.get('incomplete_details')}")
        parts: list[str] = []
        for item in response.get("output", []):
            for content in item.get("content", []):
                if content.get("type") == "output_text":
                    parts.append(content.get("text", ""))
                if content.get("type") == "refusal":
                    raise PipelineError(f"模型拒绝处理：{content.get('refusal', '')}")
        if not parts and response.get("output_text"):
            parts.append(response["output_text"])
        if not parts:
            raise PipelineError("模型响应中没有文本内容。")
        return json.loads("".join(parts))

    def transcribe(self, path: Path, *, model: str, prompt: str = "") -> dict[str, Any]:
        boundary = "----CourseNotes" + uuid.uuid4().hex
        fields = {"model": model}
        if prompt:
            fields["prompt"] = prompt[:4000]
        body = bytearray()
        for name, value in fields.items():
            body.extend(f"--{boundary}\r\nContent-Disposition: form-data; name=\"{name}\"\r\n\r\n{value}\r\n".encode())
        body.extend(
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{path.name}\"\r\n"
            "Content-Type: audio/mpeg\r\n\r\n".encode()
        )
        body.extend(path.read_bytes())
        body.extend(f"\r\n--{boundary}--\r\n".encode())
        request = urllib.request.Request(
            self.base_url + "/audio/transcriptions",
            data=bytes(body),
            headers={
                "Authorization": f"Bearer {self.api_key}",
                "Content-Type": f"multipart/form-data; boundary={boundary}",
            },
            method="POST",
        )
        try:
            with urllib.request.urlopen(request, timeout=900) as response:
                return json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")
            raise PipelineError(f"转录 API 返回 {exc.code}：{detail[:1500]}") from exc


NOTE_SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {
        "title": {"type": "string"},
        "markdown": {"type": "string"},
        "key_terms": {"type": "array", "items": {"type": "string"}},
        "uncertainties": {"type": "array", "items": {"type": "string"}},
        "source_locators": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["title", "markdown", "key_terms", "uncertainties", "source_locators"],
    "additionalProperties": False,
}


QA_SCHEMA: dict[str, Any] = {
    "type": "object",
    "properties": {
        "summary": {"type": "string"},
        "unsupported_claims": {"type": "array", "items": {"type": "string"}},
        "missing_topics": {"type": "array", "items": {"type": "string"}},
        "inconsistencies": {"type": "array", "items": {"type": "string"}},
        "manual_checks": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["summary", "unsupported_claims", "missing_topics", "inconsistencies", "manual_checks"],
    "additionalProperties": False,
}


PROFILE_RULES = {
    "stem": "重点还原定义、假设、定理、公式、符号、推导、算法、例题、边界条件与反例。区分直觉说明和严格证明。",
    "economics": "重点解释概念、模型假设、因果机制、均衡条件、公式的经济含义、实证与规范判断、模型边界。",
    "humanities": "重点还原问题意识、时间线、人物、文本证据、论证结构、解释分歧，并区分事实、作者观点与教师评价。",
    "law": "重点整理法条来源、构成要件、法律效果、原则与例外、案例争点、裁判理由、不同学说及适用法域和时间。",
    "language": "重点整理生词、词性、搭配、语域、语法规则、例外、正反例、发音和教师纠正的常见错误。",
    "general": "根据材料类型识别定义、论点、证据、机制、例子、争议和结论，保留教师特别强调的内容。",
}


BASE_INSTRUCTIONS = """你是严谨的课程笔记编辑。只能把给定材料作为教师讲授内容的证据；外部常识若确有必要，必须明确写成“补充说明”。不得编造引用、例子、课程要求或教师观点。

输出简体中文 Markdown 正文。恢复讲解逻辑，不要逐句摘要。首次出现的术语需要准确解释，重要英文术语可在括号中保留。每个公式必须使用 LaTeX，并解释所有符号、成立条件、直觉与用途。尽量使用完整段落，少用加粗与列表，使用中文标点。识别错误若无法确定，写成〔存疑：……〕。正文标题从二级标题开始。"""


def make_chunks(evidence: list[Evidence], maximum_chars: int = 18000) -> list[list[Evidence]]:
    chunks: list[list[Evidence]] = []
    current: list[Evidence] = []
    size = 0
    for item in evidence:
        if current and size + len(item.text) > maximum_chars:
            chunks.append(current)
            current = []
            size = 0
        if len(item.text) <= maximum_chars:
            current.append(item)
            size += len(item.text)
        else:
            for start in range(0, len(item.text), maximum_chars):
                if current:
                    chunks.append(current)
                    current = []
                    size = 0
                piece = item.text[start : start + maximum_chars]
                chunks.append([Evidence(item.source, item.locator, piece, item.kind)])
    if current:
        chunks.append(current)
    return chunks


def evidence_as_text(items: Iterable[Evidence]) -> str:
    sections = []
    for item in items:
        sections.append(f"[来源：{item.source}；位置：{item.locator}；类型：{item.kind}]\n{item.text}")
    return "\n\n".join(sections)


def render_qa(data: dict[str, Any]) -> str:
    lines = ["# 自动质检报告", "", data.get("summary", ""), ""]
    labels = [
        ("无法由材料直接支持的陈述", "unsupported_claims"),
        ("可能遗漏的主题", "missing_topics"),
        ("术语、公式或论述不一致", "inconsistencies"),
        ("建议人工核对", "manual_checks"),
    ]
    for title, key in labels:
        lines.extend([f"## {title}", ""])
        values = data.get(key) or []
        lines.extend([f"- {value}" for value in values] or ["未发现。"])
        lines.append("")
    return "\n".join(lines).rstrip() + "\n"


class CoursePipeline:
    def __init__(self, course_dir: Path):
        self.course_dir = course_dir.resolve()
        config_path = self.course_dir / "course.toml"
        if not config_path.exists():
            raise PipelineError(f"找不到课程配置：{config_path}")
        self.config = read_toml(config_path)

    def build(self, lecture_dir: Path, *, offline: bool = False, force: bool = False) -> Path:
        lecture_dir = lecture_dir.resolve()
        source_dir = lecture_dir / "source"
        if not source_dir.exists():
            source_dir = lecture_dir
        extracted_dir = lecture_dir / "extracted"
        generated_dir = lecture_dir / "generated"
        final_dir = lecture_dir / "final"
        extracted_dir.mkdir(parents=True, exist_ok=True)
        generated_dir.mkdir(parents=True, exist_ok=True)
        final_dir.mkdir(parents=True, exist_ok=True)

        files = [
            path for path in sorted(source_dir.rglob("*"))
            if path.is_file() and path.name not in IGNORED_NAMES and path.suffix.lower() in MEDIA_EXTENSIONS | DOCUMENT_EXTENSIONS
        ]
        references_path = source_dir / ".external-sources.json"
        if references_path.exists():
            try:
                references = json.loads(references_path.read_text(encoding="utf-8"))
                files.extend(
                    path for value in references
                    if (path := Path(value)).is_file() and path.suffix.lower() in MEDIA_EXTENSIONS | DOCUMENT_EXTENSIONS
                )
            except (json.JSONDecodeError, OSError):
                raise PipelineError(f"外部材料索引损坏：{references_path}")
        files = list(dict.fromkeys(path.resolve() for path in files))
        if not files:
            raise PipelineError(f"没有在 {source_dir} 找到可处理的课程材料。")
        manifest = {}
        for path in files:
            try:
                key = str(path.relative_to(source_dir))
            except ValueError:
                key = str(path)
            manifest[key] = file_hash(path)
        state_path = lecture_dir / ".course-notes-state.json"
        previous = json.loads(state_path.read_text(encoding="utf-8")) if state_path.exists() else {}
        config_digest = file_hash(self.course_dir / "course.toml")
        signature = hashlib.sha256(json.dumps([manifest, config_digest], sort_keys=True).encode()).hexdigest()
        draft_path = generated_dir / "draft.md"
        if (
            not force
            and previous.get("signature") == signature
            and previous.get("offline") is offline
            and draft_path.exists()
        ):
            return draft_path

        provider = None if offline else OpenAIProvider(base_url=self.config.get("models", {}).get("base_url") or None)
        evidence: list[Evidence] = []
        image_dir = extracted_dir / "slide_images"
        for path in files:
            if path.suffix.lower() in DOCUMENT_EXTENSIONS:
                evidence.extend(extract_document(path, image_dir))
            elif path.suffix.lower() in MEDIA_EXTENSIONS:
                transcript_file = extracted_dir / f"{path.stem}.transcript.json"
                if transcript_file.exists() and not force:
                    transcript = json.loads(transcript_file.read_text(encoding="utf-8"))
                elif provider is not None:
                    chunks = split_audio(path, extracted_dir / f"{path.stem}_audio")
                    model = self.config.get("models", {}).get("transcriber", "gpt-transcribe")
                    terms = ", ".join(self.config.get("course", {}).get("terms", []))
                    transcript_chunks = []
                    offset = 0.0
                    for chunk in chunks:
                        result = provider.transcribe(chunk, model=model, prompt=terms)
                        duration = media_duration(chunk)
                        transcript_chunks.append({
                            "start": offset,
                            "end": offset + duration,
                            "text": result.get("text", ""),
                            "languages": result.get("languages", []),
                        })
                        offset += duration
                    transcript = {"source": path.name, "chunks": transcript_chunks}
                    write_json(transcript_file, transcript)
                else:
                    evidence.append(Evidence(path.name, "待转录", "离线模式未转录此媒体文件。", "media"))
                    continue
                for part in transcript.get("chunks", []):
                    locator = f"{seconds_to_clock(part.get('start', 0))}–{seconds_to_clock(part.get('end', 0))}"
                    evidence.append(Evidence(path.name, locator, part.get("text", ""), "transcript"))

        evidence = [item for item in evidence if item.text.strip()]
        write_json(extracted_dir / "evidence.json", [item.as_dict() for item in evidence])
        atomic_write(extracted_dir / "evidence.md", evidence_as_text(evidence) + "\n")

        if offline:
            lines = [f"# {lecture_dir.name}：证据包", "", "当前为离线构建，以下内容尚未由模型整理。", ""]
            lines.extend([f"## {item.source}｜{item.locator}\n\n{item.text}\n" for item in evidence])
            atomic_write(draft_path, "\n".join(lines))
            atomic_write(generated_dir / "qa_report.md", "# 自动质检报告\n\n离线构建未运行语义质检。\n")
        else:
            assert provider is not None
            profile = self.config.get("course", {}).get("profile", "general")
            profile_rule = PROFILE_RULES.get(profile, PROFILE_RULES["general"])
            model = self.config.get("models", {}).get("writer", "gpt-5-mini")
            custom_prompt_path = self.course_dir / "prompt.md"
            custom = custom_prompt_path.read_text(encoding="utf-8") if custom_prompt_path.exists() else ""
            section_dir = generated_dir / "sections"
            section_dir.mkdir(parents=True, exist_ok=True)
            sections: list[dict[str, Any]] = []
            for index, chunk in enumerate(make_chunks(evidence), 1):
                cache = section_dir / f"{index:03d}.json"
                if cache.exists() and not force:
                    section = json.loads(cache.read_text(encoding="utf-8"))
                else:
                    user_input = (
                        f"课程：{self.config.get('course', {}).get('title', self.course_dir.name)}\n"
                        f"当前材料批次：{index}\n\n{evidence_as_text(chunk)}"
                    )
                    instructions = BASE_INSTRUCTIONS + "\n\n学科规则：" + profile_rule
                    if custom:
                        instructions += "\n\n课程自定义规则：\n" + custom
                    section = provider.structured_response(
                        model=model,
                        instructions=instructions,
                        user_input=user_input,
                        schema_name="course_note_section",
                        schema=NOTE_SCHEMA,
                    )
                    write_json(cache, section)
                sections.append(section)
            title = self.config.get("course", {}).get("title", self.course_dir.name)
            body = [f"# {title}｜{lecture_dir.name}", ""]
            for section in sections:
                markdown = section.get("markdown", "").strip()
                body.extend([markdown, ""])
            body.append("---\n")
            body.append("本笔记由课程材料自动整理；存疑项与材料覆盖情况见同目录下的 `qa_report.md`。\n")
            atomic_write(draft_path, "\n".join(body))

            complete_evidence = evidence_as_text(evidence)
            qa_input = f"【原始证据】\n{complete_evidence[:50000]}\n\n【生成笔记】\n{draft_path.read_text(encoding='utf-8')[:50000]}"
            reviewer = self.config.get("models", {}).get("reviewer", model)
            qa = provider.structured_response(
                model=reviewer,
                instructions="你是课程笔记审计员。只做证据覆盖、无依据陈述、术语一致性和人工复核点检查；不要重写笔记。",
                user_input=qa_input,
                schema_name="course_note_audit",
                schema=QA_SCHEMA,
            )
            write_json(generated_dir / "qa_report.json", qa)
            atomic_write(generated_dir / "qa_report.md", render_qa(qa))

        write_json(state_path, {"signature": signature, "manifest": manifest, "completed_at": time.time(), "offline": offline})
        return draft_path


def init_course(path: Path, title: str, profile: str, downloader_root: str = "") -> Path:
    path.mkdir(parents=True, exist_ok=True)
    (path / "lectures").mkdir(exist_ok=True)
    escaped_title = title.replace('"', '\\"')
    escaped_root = downloader_root.replace("\\", "\\\\").replace('"', '\\"')
    config = f'''[course]
title = "{escaped_title}"
profile = "{profile}"
language = "zh-CN"
terms = []

[models]
writer = "gpt-5-mini"
reviewer = "gpt-5-mini"
transcriber = "gpt-transcribe"

[automation]
downloader_root = "{escaped_root}"
settle_seconds = 30
'''
    atomic_write(path / "course.toml", config)
    prompt = """请保留教师的核心 insight、例子、反例、口头补充和考试提醒。课件与讲稿冲突时不要擅自消除冲突，而要标记出来。"""
    atomic_write(path / "prompt.md", prompt + "\n")
    return path


def import_lecture(
    course_dir: Path,
    title: str,
    sources: list[Path],
    *,
    move: bool = False,
    link: bool = False,
) -> Path:
    lecture_dir = course_dir / "lectures" / slugify(title)
    source_dir = lecture_dir / "source"
    source_dir.mkdir(parents=True, exist_ok=True)
    references_path = source_dir / ".external-sources.json"
    references: list[str] = []
    if references_path.exists():
        try:
            references = json.loads(references_path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            references = []
    for source in sources:
        if not source.exists() or not source.is_file():
            raise PipelineError(f"找不到输入文件：{source}")
        destination = source_dir / source.name
        if destination.resolve() == source.resolve():
            continue
        if move:
            shutil.move(str(source), destination)
        elif link:
            if source.drive.casefold() == destination.drive.casefold():
                if destination.exists():
                    destination.unlink()
                try:
                    os.link(source, destination)
                except OSError:
                    resolved = str(source.resolve())
                    if resolved not in references:
                        references.append(resolved)
            else:
                resolved = str(source.resolve())
                if resolved not in references:
                    references.append(resolved)
        else:
            shutil.copy2(source, destination)
    if references:
        write_json(references_path, references)
    return lecture_dir


def downloader_output_root(downloader_root: Path) -> Path:
    app_root = downloader_root / "pku_recorder_app"
    settings_path = app_root / "settings.json"
    if settings_path.exists():
        try:
            settings = json.loads(settings_path.read_text(encoding="utf-8"))
            configured = settings.get("download_dir", "").strip()
            if configured:
                return Path(configured)
        except (json.JSONDecodeError, OSError):
            pass
    return app_root / "downloads"


def discover_downloader_materials(downloader_root: Path, course_name: str = "") -> list[Path]:
    app_root = downloader_root / "pku_recorder_app"
    needle = course_name.casefold().strip()
    job_files: list[Path] = []
    for instance in [app_root, app_root / "dist" / "PKU-Recorder"]:
        settings_path = instance / "settings.json"
        output_root = instance / "downloads"
        if settings_path.exists():
            try:
                configured = json.loads(settings_path.read_text(encoding="utf-8")).get("download_dir", "").strip()
                if configured:
                    output_root = Path(configured)
            except (json.JSONDecodeError, OSError):
                pass
        jobs_path = instance / "jobs.json"
        if jobs_path.exists():
            try:
                jobs = json.loads(jobs_path.read_text(encoding="utf-8")).get("jobs", {})
                for job in jobs.values():
                    if job.get("status") != "done" or not job.get("file"):
                        continue
                    if needle and needle not in job.get("course_name", "").casefold():
                        continue
                    candidate = output_root / Path(job["file"])
                    if candidate.is_file() and candidate.suffix.lower() in MEDIA_EXTENSIONS | DOCUMENT_EXTENSIONS:
                        job_files.append(candidate)
            except (json.JSONDecodeError, OSError, AttributeError):
                pass
    roots = [downloader_output_root(downloader_root), app_root / "downloads", app_root / "dist" / "PKU-Recorder" / "downloads"]
    packaged_settings = app_root / "dist" / "PKU-Recorder" / "settings.json"
    if packaged_settings.exists():
        try:
            value = json.loads(packaged_settings.read_text(encoding="utf-8")).get("download_dir", "").strip()
            if value:
                roots.append(Path(value))
        except (json.JSONDecodeError, OSError):
            pass
    roots = list(dict.fromkeys(path.resolve() for path in roots))
    search_roots: list[Path] = []
    for root in roots:
        if not root.exists():
            continue
        if not needle:
            search_roots.append(root)
            continue
        try:
            first_level = list(root.iterdir())
        except OSError:
            continue
        search_roots.extend(path for path in first_level if needle in path.name.casefold())
        if not any(needle in path.name.casefold() for path in first_level):
            for parent in first_level:
                if not parent.is_dir():
                    continue
                try:
                    search_roots.extend(path for path in parent.iterdir() if needle in path.name.casefold())
                except OSError:
                    continue
    if not search_roots:
        return list(dict.fromkeys(path.resolve() for path in job_files))
    discovered: list[Path] = list(job_files)
    for search_root in search_roots:
        if search_root.is_file():
            if search_root.suffix.lower() in MEDIA_EXTENSIONS | DOCUMENT_EXTENSIONS:
                discovered.append(search_root)
            continue
        discovered.extend(
            path for path in search_root.rglob("*")
            if path.is_file() and path.suffix.lower() in MEDIA_EXTENSIONS | DOCUMENT_EXTENSIONS
        )
    return list(dict.fromkeys(path.resolve() for path in discovered))

